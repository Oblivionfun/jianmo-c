from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
root=Path(__file__).resolve().parent; d=Document();
for sec in d.sections:
 sec.top_margin=sec.bottom_margin=Pt(72); sec.left_margin=sec.right_margin=Pt(90)
def add(t,b=False):
 p=d.add_paragraph(); r=p.add_run(t); r.font.name='Arial Unicode MS'; r.font.size=Pt(11); r.font.bold=b; rf=r._element.get_or_add_rPr().get_or_add_rFonts(); rf.set(qn('w:eastAsia'),'Arial Unicode MS'); return p
add('2026 CUMCM C题 AI工具使用详情',True)
add('工具与模型：Codex（GPT-6 Astra/工作区代理）；本地Python环境（numpy、pandas、scipy、openpyxl、matplotlib、python-docx）。')
add('使用目的：检索微网储能调度与滚动优化资料；辅助整理题面与数据审计；编写和调试线性规划/MILP与滚动求解代码；生成图形、结果表和论文排版。')
add('代表性提示词：请按CUMCM本科组约束分析C题；统一能量平衡与SOC符号；实现0/6/12/18滚动更新并锁定已执行前缀；核对结果表、图形和论文中的数值一致性。')
add('采用内容：代码框架、公式排版建议、文献检索线索、结果审计清单和文字初稿。')
add('人工修改与验证：参赛队员核对题面、输入文件、单位和信息集；修正充放电SOC符号、Q3滚动实现和Q4-3独立复算；运行求解器、结果校验、输入XML审计、图形审计、Word OOXML校验和最终文稿复核。')
add('责任边界：AI未替代参赛队员的建模判断、参数选择、结果解释和最终提交责任。')
d.save(root/'AI工具使用详情.docx')
