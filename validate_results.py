from pathlib import Path
import json, numpy as np
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent; OUT=ROOT/'results'; FIG=ROOT/'figures'
summary=json.loads((OUT/'summary.json').read_text())
q1=summary['q1']; assert abs(q1['terminal_soc']-6000)<1e-6
assert 1200-1e-6<=q1['soc_min']<=10800+1e-6 and 1200-1e-6<=q1['soc_max']<=10800+1e-6
assert q1['purchase_kwh']>0
for fn,sheets in {'result1.xlsx':['计划购电量','充放电量'],'result2.xlsx':['计划购电量','充放电量','紧急购电量'],'result3.xlsx':['计划购电量','调整购电量','充放电量','紧急购电量'],'result4-2.xlsx':['计划购电量','充放电量','紧急购电量'],'result4-3.xlsx':['计划购电量','调整购电量','充放电量','紧急购电量']}.items():
 p=OUT/fn; assert p.exists() and p.stat().st_size>0
 wb=load_workbook(p,read_only=True,data_only=True); assert wb.sheetnames==sheets,(fn,wb.sheetnames)
 for ws in wb.worksheets: assert ws.max_row>=2 and ws.max_column>=2
# required figure coverage: 3 categories x q1..q4 (logical pngs)
for q in range(1,5):
 for cat in ('raw','process','result'):
  found=list(FIG.glob(f'{cat}_q{q}_*.png'))
  assert found,(cat,q)
print('PASS result validation')
print('Q1 purchase_kwh',q1['purchase_kwh'],'Q1 soc',q1['soc_min'],q1['soc_max'])
print('aggregates',summary['aggregate'])
print('fig_png',len(list(FIG.glob('*.png'))),'fig_svg',len(list(FIG.glob('*.svg'))))
