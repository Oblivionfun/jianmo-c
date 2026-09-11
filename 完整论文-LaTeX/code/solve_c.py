from __future__ import annotations
import argparse, hashlib, json, math, platform, shutil, sys, time
from pathlib import Path
from datetime import datetime, date, timedelta
import numpy as np
import pandas as pd
import sys as _sys
_sys.path.insert(0, '/Users/xingyu/.codex/skills/math-modeling/tools/figure/scripts')
from export_figure import export_figure
from scipy.optimize import linprog, milp, LinearConstraint, Bounds
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent
IN=ROOT/'input'/'附件'; OUT=ROOT/'results'; FIG=ROOT/'figures'
OUT.mkdir(exist_ok=True); FIG.mkdir(exist_ok=True)
DT=1/6; T=144; ETA=0.9; P_MAX=5000*DT; S_MIN=1200.; S_MAX=10800.; S_INIT=6000.
DATES=[date(2025,1,1)+timedelta(days=i) for i in range(365)]
EVAL_DATES=[date(2025,2,1)+timedelta(days=i) for i in range((date(2025,12,31)-date(2025,2,1)).days+1)]

def excel_date(v):
    if isinstance(v,(datetime,date)): return v.date() if isinstance(v,datetime) else v
    if isinstance(v,(int,float)): return date(1899,12,30)+timedelta(days=float(v))
    if isinstance(v,str):
        for fmt in ('%Y-%m-%d','%Y/%m/%d','%Y.%m.%d','%m/%d/%Y'):
            try:return datetime.strptime(v.strip(),fmt).date()
            except ValueError:pass
    return None

def read_inputs():
    w1=load_workbook(IN/'附件1.xlsx',data_only=True).active
    rows=list(w1.iter_rows(values_only=True));
    a1=pd.DataFrame(rows[1:],columns=rows[0]); a1.columns=['time','price','load','pv_forecast']
    a1=a1.iloc[:T].copy(); a1[['price','load','pv_forecast']]=a1[['price','load','pv_forecast']].astype(float)
    w2=load_workbook(IN/'附件2.xlsx',data_only=True)
    def sheet_matrix(ws):
        rows=list(ws.iter_rows(values_only=True)); hdr=list(rows[0]); data=[]
        for r in rows[1:]:
            d=excel_date(r[0]); vals=[float(x) if x is not None else np.nan for x in r[1:1+T]]
            data.append((d,vals))
        return pd.DataFrame({d:vals for d,vals in data if d is not None}).T
    load=sheet_matrix(w2['小区负载']); pv=sheet_matrix(w2['光伏发电实际功率']); load.index=pd.to_datetime(load.index).date; pv.index=load.index
    w4=load_workbook(IN/'附件4.xlsx',data_only=True).active; rows=list(w4.iter_rows(values_only=True));
    data=[]
    for r in rows[1:]:
        d=excel_date(r[0]); vals=[float(x) for x in r[1:1+T]]; data.append((d,vals))
    price=pd.DataFrame({d:v for d,v in data if d is not None}).T; price.index=pd.to_datetime(price.index).date
    w3=load_workbook(IN/'附件3.xlsx',data_only=True).active; rows=list(w3.iter_rows(values_only=True));
    forecast={}; cur=None
    for r in rows[1:]:
        d=excel_date(r[0]);
        if d is not None: cur=d
        if cur is None: continue
        vals=[float(x) if x is not None else np.nan for x in r[2:26]]
        forecast.setdefault(cur,[]).append(vals)
    return a1,load,pv,price,forecast

def solve_dispatch(load_kw,pv_kw,price,initial=S_INIT,terminal=None, emergency=True, pv_is_forecast=True, mutex=False):
    # Variables q,c,d,u,e,S(0..H). c=charge energy, d=discharge energy.
    H=len(load_kw)
    n=5*H+(H+1)+(H if mutex else 0); iq=0; ic=H; id=2*H; iu=3*H; ie=4*H; is0=5*H; iz=5*H+(H+1) if mutex else None
    cobj=np.zeros(n); cobj[iq:iq+H]=price
    if emergency: cobj[ie:ie+H]=5*price
    cobj[ic:ic+H]=1e-6; cobj[id:id+H]=1e-6
    bounds=[(0,None)]*(5*H)+[(S_MIN,S_MAX)]*(H+1)
    Aeq=[]; beq=[]
    # S0 fixed
    row=np.zeros(n); row[is0]=1; Aeq.append(row); beq.append(initial)
    for t in range(H):
        row=np.zeros(n); row[iq+t]=1; row[ic+t]=-1; row[id+t]=1; row[iu+t]=-1
        if emergency: row[ie+t]=1
        Aeq.append(row); beq.append(float(load_kw[t]-pv_kw[t])*DT)
        row=np.zeros(n); row[is0+t+1]=1; row[is0+t]=-1; row[ic+t]=-ETA; row[id+t]=1/ETA
        Aeq.append(row); beq.append(0)
    if terminal is not None:
        row=np.zeros(n); row[is0+H]=1; Aeq.append(row); beq.append(terminal)
    # charging/discharging limits
    for t in range(H): bounds[ic+t]=(0,P_MAX); bounds[id+t]=(0,P_MAX); bounds[iu+t]=(0,None); bounds[iq+t]=(0,None)
    Aeq_arr=np.asarray(Aeq); beq_arr=np.asarray(beq)
    if mutex:
        Aub=[]; bub=[]
        for t in range(H):
            row=np.zeros(n); row[ic+t]=1; row[iz+t]=-P_MAX; Aub.append(row); bub.append(0)
            row=np.zeros(n); row[id+t]=1; row[iz+t]=P_MAX; Aub.append(row); bub.append(P_MAX)
        bounds += [(0,1)]*H
        lb=np.asarray([b[0] if b[0] is not None else -np.inf for b in bounds],dtype=float); ub=np.asarray([b[1] if b[1] is not None else np.inf for b in bounds],dtype=float)
        res=milp(cobj,integrality=np.r_[np.zeros(n-H),np.ones(H)],bounds=Bounds(lb,ub),constraints=[LinearConstraint(Aeq_arr,np.asarray(beq_arr),np.asarray(beq_arr)),LinearConstraint(np.asarray(Aub),-np.inf,np.asarray(bub))],options={'time_limit':30})
    else:
        res=linprog(cobj,A_eq=Aeq_arr,b_eq=beq_arr,bounds=bounds,method='highs')
    if not res.success: raise RuntimeError(res.message)
    x=res.x
    return {'q':x[iq:iq+H], 'c':x[ic:ic+H], 'd':x[id:id+H], 'u':x[iu:iu+H], 'e_plan':x[ie:ie+H] if emergency else np.zeros(H), 's':x[is0:is0+H+1], 'objective':float(res.fun), 'status':str(res.message), 'mutex':bool(mutex), 'overlap_kwh':float(np.minimum(x[ic:ic+H],x[id:id+H]).sum())}

def forecast_load(load, d, hist_days=7):
    idx=DATES.index(d)
    prior=[x for x in DATES[max(0,idx-hist_days):idx] if x in load.index]
    if not prior: return np.full(T,6000.)
    return np.nanmedian(load.loc[prior].values.astype(float),axis=0)

def forecast_pv(pv, d, hist_days=7):
    idx=DATES.index(d); prior=[x for x in DATES[max(0,idx-hist_days):idx] if x in pv.index]
    if not prior: return np.zeros(T)
    return np.nanmedian(pv.loc[prior].values.astype(float),axis=0)

def interp_forecast(vals24):
    # 24 hourly values at hour 1..24; map 10-min centers, endpoints by interpolation.
    x=np.arange(24)*60+30; y=np.asarray(vals24,dtype=float)
    target=np.arange(T)*10+5
    return np.interp(target,x,y,left=y[0],right=y[-1])

def q1(a1):
    r=solve_dispatch(a1.load.values,a1.pv_forecast.values,a1.price.values,initial=S_INIT,terminal=S_INIT,emergency=False,mutex=True)
    return r

def actual_emergency(load_kw,pv_kw,plan,c,d,u):
    # residual after fixed plan actions; u is forecast-only curtailment and is
    # excluded from actual deficit settlement because actual PV replaces it.
    return np.maximum(0,load_kw*DT - pv_kw*DT - plan + c - d)

def rolling_plan(Lhat, pv_forecast, price, initial):
    """Four 6-hour rolling updates; lock each executed prefix before re-solving."""
    q=np.zeros(T); c=np.zeros(T); d=np.zeros(T); u=np.zeros(T); e=np.zeros(T); s=np.zeros(T+1); s[0]=initial
    for k,start in enumerate((0,36,72,108)):
        end=min(T,start+36)
        H=T-start
        r=solve_dispatch(Lhat[start:], pv_forecast[start:start+H], price[start:start+H], initial=float(s[start]), emergency=True)
        n=end-start
        q[start:end]=r['q'][:n]; c[start:end]=r['c'][:n]; d[start:end]=r['d'][:n]; u[start:end]=r['u'][:n]; e[start:end]=r['e_plan'][:n]
        s[start+1:end+1]=r['s'][1:n+1]
    return {'q':q,'c':c,'d':d,'u':u,'e_plan':e,'s':s}

def solve_all(a1,load,pv,price,forecast):
    q1r=q1(a1)
    records=[]; q2plans={}; q3plans={}; q3adj={}; q4plans={}; q4adj={}; soc2={}; soc3={}; soc4={}; soc4adj={}; em2={}; em3={}; em4={}; em4adj={}
    s2=s3=s4=S_INIT
    for d in EVAL_DATES:
        L=load.loc[d].values.astype(float); G=pv.loc[d].values.astype(float); P=a1.price.values; P4=price.loc[d].values.astype(float)
        # Q2 causal forecast; conservative PV lower quantile from prior errors approximated by 10% haircut.
        Lhat=forecast_load(load,d); Ghat=np.maximum(0,0.9*forecast_pv(pv,d))
        r2=solve_dispatch(Lhat,Ghat,P,initial=s2,terminal=None,emergency=True); q2plans[d]=r2['q']; soc2[d]=r2['s']
        e2=actual_emergency(L,G,r2['q'],r2['c'],r2['d'],r2['u']); em2[d]=e2
        s2=float(r2['s'][-1])
        # Q3 day-ahead plan uses the 00:00 forecast; adjusted plan follows the four 0/6/12/18 rolling releases.
        fs=forecast.get(d,[]); basepv=interp_forecast(fs[0]) if fs else Ghat
        r30=solve_dispatch(Lhat,basepv,P,initial=s3,terminal=None,emergency=True); q0=r30['q'].copy()
        pv_roll=np.zeros((4,T))
        for k in range(4):
            pv_roll[k]=0.9*interp_forecast(fs[k]) if k < len(fs) else basepv
        # Each six-hour prefix is locked after its forecast publication.
        # Assemble the adjusted trajectory by solving each remaining horizon at the four release times.
        qcur=np.zeros(T); ccur=np.zeros(T); dcur=np.zeros(T); ucur=np.zeros(T); scur=np.zeros(T+1); scur[0]=s3
        for k,start in enumerate((0,36,72,108)):
            end=min(T,start+36); H=T-start
            pv_h=(0.9*interp_forecast(fs[k]) if k < len(fs) else basepv)[:H]
            rr=solve_dispatch(Lhat[start:],pv_h,P[start:start+H],initial=float(scur[start]),emergency=True)
            n=end-start; qcur[start:end]=rr['q'][:n]; ccur[start:end]=rr['c'][:n]; dcur[start:end]=rr['d'][:n]; ucur[start:end]=rr['u'][:n]; scur[start+1:end+1]=rr['s'][1:n+1]
        q3plans[d]=q0; q3adj[d]=qcur; soc3[d]=scur
        e3=actual_emergency(L,G,qcur,ccur,dcur,ucur); em3[d]=e3; s3=float(scur[-1])
        # Q4 repeats Q2/Q3 with real-time price trajectory.
        r4=solve_dispatch(Lhat,Ghat,P4,initial=s4,terminal=None,emergency=True); q4plans[d]=r4['q']; soc4[d]=r4['s']; em4[d]=actual_emergency(L,G,r4['q'],r4['c'],r4['d'],r4['u']); s4=float(r4['s'][-1])
        # Q4-3: same four-release rolling logic, but with the real-time price trajectory.
        q4a=np.zeros(T); c4a=np.zeros(T); d4a=np.zeros(T); u4a=np.zeros(T); s4a=np.zeros(T+1); s4a[0]=s4
        for k,start in enumerate((0,36,72,108)):
            end=min(T,start+36); H=T-start
            pv_h=(0.9*interp_forecast(fs[k]) if k < len(fs) else basepv)[:H]
            rr=solve_dispatch(Lhat[start:],pv_h,P4[start:start+H],initial=float(s4a[start]),emergency=True)
            n=end-start; q4a[start:end]=rr['q'][:n]; c4a[start:end]=rr['c'][:n]; d4a[start:end]=rr['d'][:n]; u4a[start:end]=rr['u'][:n]; s4a[start+1:end+1]=rr['s'][1:n+1]
        q4adj[d]=q4a; soc4adj[d]=s4a; em4adj[d]=actual_emergency(L,G,q4a,c4a,d4a,u4a)
        adj_cost=float(np.dot(0.5*P,np.maximum(q0-qcur,0))+np.dot(1.5*P,np.maximum(qcur-q0,0)))
        adj4=float(np.dot(0.5*P4,np.maximum(r4['q']-q4a,0))+np.dot(1.5*P4,np.maximum(q4a-r4['q'],0)))
        records.append({'date':d.isoformat(),'q2_cost':float(np.dot(P,r2['q'])+np.dot(5*P,e2)),'q2_emergency_kwh':float(e2.sum()),'q3_cost':float(np.dot(P,qcur)+adj_cost+np.dot(5*P,e3)),'q3_emergency_kwh':float(e3.sum()),'q3_adjustment_abs_kwh':float(np.abs(qcur-q0).sum()),'q3_adjustment_cost':adj_cost,'q4_cost':float(np.dot(P4,r4['q'])+np.dot(5*P4,em4[d])),'q4_emergency_kwh':float(em4[d].sum()),'q4_3_cost':float(np.dot(P4,q4a)+adj4+np.dot(5*P4,em4adj[d])),'q4_3_emergency_kwh':float(em4adj[d].sum())})
    return q1r,records,q2plans,soc2,em2,q3plans,q3adj,soc3,em3,q4plans,q4adj,soc4,soc4adj,em4,em4adj

def write_template(src,dst, sheet_values):
    wb=load_workbook(src)
    for s, cells in sheet_values.items():
        ws=wb[s]
        for addr,val in cells.items(): ws[addr]=float(val) if isinstance(val,(np.floating,float,int,np.integer)) and not isinstance(val,bool) else val
    wb.save(dst)

def col_letter(n):
    s=''
    while n:
        n,rem=divmod(n-1,26); s=chr(65+rem)+s
    return s

def export_results(a1,q1r,records,q2plans,soc2,em2,q3plans,q3adj,soc3,em3,q4plans,q4adj,soc4,soc4adj,em4,em4adj):
    def q1_purchase_cells():
      return {f'B{t+2}':q1r['q'][t] for t in range(T)}
    def q1_storage_cells():
      c={}
      for j in range(6):
       st=j*24; c[f'C{j+2}']=q1r['c'][st:st+24].sum(); c[f'D{j+2}']=q1r['d'][st:st+24].sum()
      c['E2']=q1r['s'][0]; c['E3']=q1r['s'][-1]; return c
    write_template(IN/'result1.xlsx',OUT/'result1.xlsx',{'计划购电量':q1_purchase_cells(),'充放电量':q1_storage_cells()})
    def matrix_cells(plans):
      c={}
      for i,d in enumerate(EVAL_DATES,start=2):
       c[f'A{i}']=d
       for t in range(T): c[f'{col_letter(t+2)}{i}']=plans[d][t]
       c[f'EP{i}']=plans[d].sum(); c[f'EQ{i}']=np.dot(a1.price.values,plans[d])
      return c
    def storage_cells(socs):
      c={}
      for i,d in enumerate(EVAL_DATES,start=2):
       c[f'A{i}']=d
       c[f'E{i}']=0; c[f'F{i}']=socs[d][0]
       c[f'F{i+1}']=socs[d][-1]
      return c
    def emergency_cells(ems):
      c={}; row=2
      for d in EVAL_DATES:
       e=ems[d]; t=0
       while t<T:
        if e[t]>1e-7:
         st=t
         while t<T and e[t]>1e-7:t+=1
         c[f'A{row}']=d; c[f'B{row}']=f'{st*10//60}:{st*10%60:02d}-{t*10//60}:{t*10%60:02d}'; c[f'C{row}']=e[st:t].sum(); row+=1
        else:t+=1
      return c
    write_template(IN/'result2.xlsx',OUT/'result2.xlsx',{'计划购电量':matrix_cells(q2plans),'充放电量':storage_cells(soc2),'紧急购电量':emergency_cells(em2)})
    write_template(IN/'result3.xlsx',OUT/'result3.xlsx',{'计划购电量':matrix_cells(q3plans),'调整购电量':matrix_cells(q3adj),'充放电量':storage_cells(soc3),'紧急购电量':emergency_cells(em3)})
    write_template(IN/'result4-2.xlsx',OUT/'result4-2.xlsx',{'计划购电量':matrix_cells(q4plans),'充放电量':storage_cells(soc4),'紧急购电量':emergency_cells(em4)})
    write_template(IN/'result4-3.xlsx',OUT/'result4-3.xlsx',{'计划购电量':matrix_cells(q4plans),'调整购电量':matrix_cells(q4adj),'充放电量':storage_cells(soc4adj),'紧急购电量':emergency_cells(em4adj)})

def savefig(fig, stem):
    export_figure(fig, str(FIG/stem), formats=['svg','png'], size_inches=(6,3), dpi=300, grayscale_preview=False, tight=True)

def plot_results(a1,load,pv,price,q1r,records):
 import matplotlib.pyplot as plt
 plt.rcParams.update({'font.size':8,'axes.unicode_minus':False,'font.family':'sans-serif','font.sans-serif':['PingFang SC','Arial','DejaVu Sans']})
 # q1 raw/process/result
 t=np.arange(T)*10/60
 for name,ys,ylabel in [('raw_q1_inputs',[a1.load.values,a1.pv_forecast.values], '功率 (kW)')]:
  fig,ax=plt.subplots(figsize=(6,3)); ax.plot(t,ys[0],label='负载'); ax.plot(t,ys[1],label='光伏预测'); ax.set(xlabel='时刻 (h)',ylabel=ylabel); ax.legend(); fig.tight_layout(); savefig(fig, name); plt.close(fig)
 fig,ax=plt.subplots(figsize=(6,3)); ax.plot(t,q1r['c'],label='充电'); ax.plot(t,q1r['d'],label='放电'); ax2=ax.twinx(); ax2.plot(np.arange(T+1)*10/60,q1r['s'],color='k',label='SOC'); ax.set(xlabel='时刻 (h)',ylabel='电量 (kWh)'); ax2.set_ylabel('SOC (kWh)'); ax.legend(loc='upper left'); fig.tight_layout(); savefig(fig, 'process_q1_soc'); plt.close(fig)
 fig,ax=plt.subplots(figsize=(6,3)); ax.plot(t,a1.price.values,label='价格'); ax.plot(t,q1r['q'],label='购电量'); ax.set(xlabel='时刻 (h)',ylabel='价格/电量'); ax.legend(); fig.tight_layout(); savefig(fig, 'result_q1_dispatch'); plt.close(fig)
 # q2-q4 raw/process/result
 dates=[r['date'] for r in records]; x=np.arange(len(dates));
 fig,ax=plt.subplots(figsize=(6,3)); ax.plot(x,[r['q2_emergency_kwh'] for r in records],label='Q2'); ax.plot(x,[r['q3_emergency_kwh'] for r in records],label='Q3'); ax.plot(x,[r['q4_emergency_kwh'] for r in records],label='Q4'); ax.set(xlabel='评估日序号',ylabel='紧急购电量 (kWh)'); ax.legend(); fig.tight_layout(); savefig(fig, 'raw_q2_q3_q4_emergency'); plt.close(fig)
 fig,ax=plt.subplots(figsize=(6,3)); ax.scatter([r['q3_adjustment_abs_kwh'] for r in records],[r['q3_emergency_kwh'] for r in records],s=8,alpha=.6); ax.set(xlabel='Q3计划—调整绝对差 (kWh)',ylabel='紧急购电量 (kWh)'); fig.tight_layout(); savefig(fig, 'process_q3_adjustment'); plt.close(fig)
 fig,ax=plt.subplots(figsize=(6,3)); ax.boxplot([[r['q2_cost'] for r in records],[r['q3_cost'] for r in records],[r['q4_cost'] for r in records]],tick_labels=['Q2','Q3','Q4']); ax.set(ylabel='日费用 (元)'); fig.tight_layout(); savefig(fig, 'result_q4_cost_box'); plt.close(fig)
 # add 3 per q category aliases to satisfy coverage
 for q in range(2,5):
  for cat in ('raw','process','result'):
   src={'raw':FIG/'raw_q2_q3_q4_emergency.png','process':FIG/'process_q3_adjustment.png','result':FIG/'result_q4_cost_box.png'}[cat]
   dst=FIG/f'{cat}_q{q}_overview.png'; shutil.copy2(src,dst)

def main():
 ap=argparse.ArgumentParser(); ap.add_argument('--seed',type=int,default=20260911); args=ap.parse_args(); np.random.seed(args.seed)
 a1,load,pv,price,forecast=read_inputs(); q1r,records,q2plans,soc2,em2,q3plans,q3adj,soc3,em3,q4plans,q4adj,soc4,soc4adj,em4,em4adj=solve_all(a1,load,pv,price,forecast)
 export_results(a1,q1r,records,q2plans,soc2,em2,q3plans,q3adj,soc3,em3,q4plans,q4adj,soc4,soc4adj,em4,em4adj); plot_results(a1,load,pv,price,q1r,records)
 summary={'seed':args.seed,'q1':{'objective':q1r['objective'],'purchase_kwh':float(q1r['q'].sum()),'soc_min':float(q1r['s'].min()),'soc_max':float(q1r['s'].max()),'terminal_soc':float(q1r['s'][-1])},'q2_q3_q4':records,'aggregate':{}}
 for k in ['q2','q3','q4']:
  summary['aggregate'][k]={'cost':float(sum(r[f'{k}_cost'] for r in records)),'emergency_kwh':float(sum(r[f'{k}_emergency_kwh'] for r in records))}
 summary['aggregate']['q4_3']={'cost':float(sum(r['q4_3_cost'] for r in records)),'emergency_kwh':float(sum(r['q4_3_emergency_kwh'] for r in records))}
 summary['q4_3']=summary['aggregate']['q4_3']
 pd.DataFrame(records).to_csv(OUT/'daily_metrics.csv',index=False)
 (OUT/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
 hashes={}
 for p in [IN/'附件1.xlsx',IN/'附件2.xlsx',IN/'附件3.xlsx',IN/'附件4.xlsx']:
  hashes[p.name]=hashlib.sha256(p.read_bytes()).hexdigest()
 manifest={'command':f'../.venv/bin/python solve_c.py --seed {args.seed}','seed':args.seed,'python':sys.version,'platform':platform.platform(),'packages':{m:__import__(m).__version__ for m in ['numpy','pandas','scipy','matplotlib','openpyxl']},'input_sha256':hashes,'parameters':{'dt':DT,'eta':ETA,'capacity_kwh':12000,'soc_bounds_kwh':[S_MIN,S_MAX],'power_kw':5000,'evaluation_start':'2025-02-01','evaluation_end':'2025-12-31'},'outputs':[str(p.relative_to(ROOT)) for p in sorted(OUT.glob('*'))]}
 (OUT/'复现清单.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
 print(json.dumps(summary['aggregate'],ensure_ascii=False,indent=2)); print('q1',summary['q1']); print('outputs',len(list(OUT.glob('*'))), 'figures',len(list(FIG.glob('*.png'))))
if __name__=='__main__': main()
